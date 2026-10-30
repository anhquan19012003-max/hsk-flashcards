"""Export a level to review/hsk{N}_review.xlsx for Quan to edit.

Usage: python scripts/export_review.py N [--only-flagged] [--force]
Refuses to overwrite an existing review file (it may hold Quan's edits) unless --force is given.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
HEADERS = ["id", "Chữ Hán", "Pinyin", "Nghĩa", "Từ loại", "Hán Việt", "Ví dụ (zh)", "Ví dụ (pinyin)",
           "Ví dụ (vi)", "Cờ kiểm tra", "Cần kiểm Hán Việt", "Ghi chú / sửa"]
WIDTHS = [9, 12, 14, 30, 16, 16, 28, 36, 36, 36, 12, 30]


def main():
    args = sys.argv[1:]
    level = int(next(a for a in args if a.isdigit()))
    only_flagged = "--only-flagged" in args
    data = json.loads((ROOT / "data" / "enriched" / f"hsk{level}.json").read_text(encoding="utf-8"))
    out = ROOT / "review" / f"hsk{level}_review.xlsx"
    out.parent.mkdir(exist_ok=True)
    if out.exists() and "--force" not in args:
        sys.exit(f"{out} exists and may contain review edits; use --force to overwrite or move it first.")
    wb = Workbook()
    ws = wb.active
    ws.title = f"HSK{level}"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    n = 0
    for w in data["words"]:
        flags = w.get("flags", [])
        needs = bool(flags) or w.get("hanviet_check")
        if only_flagged and not needs:
            continue
        ex = w.get("example") or {}
        ws.append([w["id"], w["hanzi"], w["pinyin"], w["meaning"], w["pos"], w.get("hanviet", ""),
                   ex.get("zh", ""), ex.get("pinyin", ""), ex.get("vi", ""),
                   "; ".join(flags + ([w["hanviet_note"]] if w.get("hanviet_note") else [])),
                   "x" if w.get("hanviet_check") else "", ""])
        if needs:
            for c in ws[ws.max_row]:
                c.fill = yellow
        n += 1
    for i, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    print(f"{out}: {n} rows")


if __name__ == "__main__":
    main()
